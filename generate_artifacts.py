"""
generate_artifacts.py
---------------------
Script to build documentation/TDMA_Project_Report.pdf and presentation/TDMA_Project_Presentation.pptx
"""

import os
import re
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, HRFlowable
from reportlab.lib import colors
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


def clean_markdown_to_reportlab_html(text: str) -> str:
    # Replace & < > first
    t = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    # Convert **bold** to <b>bold</b> using regex
    t = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', t)
    # Convert *italic* to <i>italic</i> using regex
    t = re.sub(r'\*(.*?)\*', r'<i>\1</i>', t)
    return t


def generate_pdf():
    pdf_path = os.path.join("documentation", "TDMA_Project_Report.pdf")
    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=letter,
        rightMargin=54,
        leftMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=22,
        leading=26,
        textColor=colors.HexColor('#1a365d'),
        spaceAfter=6
    )

    subtitle_style = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#4a5568'),
        spaceAfter=15
    )

    h1_style = ParagraphStyle(
        'Heading1Custom',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#2b6cb0'),
        spaceBefore=14,
        spaceAfter=6
    )

    body_style = ParagraphStyle(
        'BodyCustom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#2d3748'),
        spaceAfter=8
    )

    code_style = ParagraphStyle(
        'CodeCustom',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#1a202c'),
        backColor=colors.HexColor('#f7fafc'),
        spaceAfter=8,
        leftIndent=10
    )

    story = []

    story.append(Paragraph("TDMA Schedule Planner and Optimizer for Wireless Networks", title_style))
    story.append(Paragraph("Technical Project Report — Wireless Protocol Development Role", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#2b6cb0'), spaceAfter=15))

    md_path = os.path.join("documentation", "TDMA_Project_Report.md")
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []

    for line in lines:
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code_block:
                code_text = "<br/>".join(code_lines).replace(" ", "&nbsp;")
                story.append(Paragraph(code_text, code_style))
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
            continue

        if in_code_block:
            safe_line = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").rstrip()
            code_lines.append(safe_line)
            continue

        if stripped.startswith("# "):
            continue
        elif stripped.startswith("## "):
            h_text = clean_markdown_to_reportlab_html(stripped[3:].strip())
            story.append(Paragraph(h_text, h1_style))
        elif stripped.startswith("### "):
            h_text = clean_markdown_to_reportlab_html(stripped[4:].strip())
            story.append(Paragraph(h_text, h1_style))
        elif stripped == "---":
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#cbd5e0'), spaceBefore=8, spaceAfter=8))
        elif stripped:
            p_text = clean_markdown_to_reportlab_html(stripped)
            story.append(Paragraph(p_text, body_style))

    doc.build(story)
    print(f"[SUCCESS] PDF generated at {pdf_path}")


def generate_pptx():
    pptx_path = os.path.join("presentation", "TDMA_Project_Presentation.pptx")
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    slides_data = [
        {
            "title": "TDMA Schedule Planner and Optimizer",
            "subtitle": "Wireless Protocol Development — Technical Presentation",
            "bullets": [
                "Collision-Free Spatial Reuse TDMA Scheduling algorithm",
                "Formulated as Distance-2 Graph Coloring on Communication Graph G",
                "DSATUR Heuristic + Greedy Local Search Post-Optimization Pass",
                "Official EMANE TDMA Radio Model Integration (1 ms slot timing)",
                "Full Automated Test Suite: 113 / 113 Tests Passing (0.29s)"
            ]
        },
        {
            "title": "Problem Statement & Interference Constraints",
            "subtitle": "Why Ordinary Graph Coloring is Insufficient",
            "bullets": [
                "Direct Interference (1-Hop): Adjacent nodes within radio range cannot share a slot.",
                "Hidden Terminal Interference (2-Hop): Common 1-hop neighbors cause collisions at receiver.",
                "Ordinary Graph Coloring solves 1-hop conflicts only — leaves hidden terminals unaddressed.",
                "Distance-2 Graph Coloring is required: Edges added between all nodes within 2 hops in G."
            ]
        },
        {
            "title": "Mathematical Formulation & NP-Hardness",
            "subtitle": "Graph Theory Formulation",
            "bullets": [
                "Communication Graph G = (V, E_G): Edge exists if Euclidean_distance(u, v) <= R_radio.",
                "Conflict Graph C = (V, E_C): Edge exists if dist_G(u, v) <= 2.",
                "Decision Problem: Find chromatic number chi(C) representing minimum required frame slots.",
                "Complexity: Distance-2 Coloring is NP-hard — justifies using dynamic DSATUR heuristic."
            ]
        },
        {
            "title": "Scheduling Algorithm: DSATUR Heuristic",
            "subtitle": "Degree of Saturation (Brélaz 1979)",
            "bullets": [
                "Saturation Degree sat(v): Number of distinct colors assigned to conflict neighbors in C.",
                "Dynamic Node Selection: Select uncolored node with highest sat(v) (tie-breaker: degree in C).",
                "First-Fit Slot Assignment: Assign lowest-indexed slot s >= 0 not used by conflict neighbors.",
                "Dynamic Updates: Saturation values updated dynamically after every node assignment."
            ]
        },
        {
            "title": "Local Optimization & Spatial Reuse",
            "subtitle": "Schedule Frame Compaction",
            "bullets": [
                "Spatial Reuse: Spatially separated nodes (dist_G > 2) share the same TDMA slot.",
                "Greedy Local Search Pass: Iterates nodes deterministically, shifting to lower slots if feasible.",
                "9-slot conflict-free heuristic schedule: Compacted 16-node grid frame into 9 slots.",
                "Contiguous Normalization: Ensures slot indices range continuously from 0 to Frame_Length - 1."
            ]
        },
        {
            "title": "Part 1 Experimental Results (16-Node Grid)",
            "subtitle": "Dataset: examples/nodes16.json (Radio Range = 500m)",
            "bullets": [
                "Input Metrics: 16 Nodes, 42 Communication Links, 90 Distance-2 Conflict Edges.",
                "Resulting Frame: 9-slot conflict-free heuristic schedule.",
                "Slot Utilization: Slot 8 shared by Nodes 1, 10, 13, 16 (Spatial reuse verified).",
                "Validator Inspection: 0 conflicts detected across all 90 conflict edges (100% Valid)."
            ]
        },
        {
            "title": "Part 2 Architecture: EMANE XML Schedule Adapter",
            "subtitle": "Integration with EMANE TDMA Radio Model",
            "bullets": [
                "Schedule Adapter (schedule_adapter.py): Converts ScheduleResult into official EMANE XML.",
                "NEM ID Mapping: Deterministic alphabetical mapping (Node_01 -> NEM 1).",
                "1 ms Slot Timing: Configures slotduration='1000' (microseconds).",
                "Official Schema: Generates <emane-tdma-schedule> with nodes='1 10 13 16' attribute."
            ]
        },
        {
            "title": "Linux Network Namespace Architecture",
            "subtitle": "Preventing Host Kernel Routing Bypass",
            "bullets": [
                "Host Routing Problem: TAP devices in root namespace cause kernel to route via loopback.",
                "Namespace Isolation Solution: Isolated Linux namespaces (node-1 .. node-16).",
                "Interface Binding: TAP interface emaneX moved inside node-X namespace.",
                "Traffic Flow: Forces all IP traffic through virtual TAP interface into EMANE MAC/PHY."
            ]
        },
        {
            "title": "EMANE Event Service & Inspection",
            "subtitle": "Multicast Schedule Injection & Monitoring",
            "bullets": [
                "Event Injection: emaneevent-tdmaschedule tdma-schedule.xml -i lo",
                "Event Type: TDMAScheduleEvent (Event ID 102) broadcast over multicast.",
                "Runtime Verification: emanesh localhost get table 1 mac scheduler.ScheduleInfoTable",
                "Traffic Testing: sudo ip netns exec node-5 iperf3 -c 10.0.0.6 -u -b 1M -t 10"
            ]
        },
        {
            "title": "Software Quality & Automated Test Suite",
            "subtitle": "113 / 113 Tests Passing (0.29s)",
            "bullets": [
                "test_coloring.py (16 tests): DSATUR heuristic & color normalization.",
                "test_conflict_graph.py (13 tests): Distance-1 & Distance-2 conflict graph builder.",
                "test_emane_adapter.py (10 tests): EMANE XML schema, NEM mapping, slot timing.",
                "test_end_to_end.py & test_topology.py & test_validator.py (74 tests): End-to-end pipeline."
            ]
        },
        {
            "title": "Implementation vs. Testing Status",
            "subtitle": "Transparent Engineering Breakdown",
            "bullets": [
                "IMPLEMENTED & UNIT-TESTED: Schedule optimizer, DSATUR, Validator, EMANE XML Adapter (113 tests).",
                "CONFIGURATION GENERATED: tdma-schedule.xml, nem.xml, platform.xml (16 NEMs), transvirtual.xml.",
                "DESIGNED FOR LINUX: Network namespace script (launch_emane.sh), EMANE daemon execution.",
                "NOT RUNTIME TESTED: Live EMANE daemon and real packet delivery (Host OS is Windows)."
            ]
        },
        {
            "title": "Conclusion & Key Takeaways",
            "subtitle": "Summary for Technical Interview Panel",
            "bullets": [
                "Mathematically rigorous Distance-2 graph coloring formulation for TDMA scheduling.",
                "9-slot conflict-free heuristic schedule achieved on 16-node grid.",
                "Official EMANE TDMA XML event schema integration with 1 ms slot timing.",
                "100% passing automated test suite (113/113) with full defensible documentation."
            ]
        }
    ]

    for s_data in slides_data:
        slide = prs.slides.add_slide(blank_layout)

        header_shape = slide.shapes.add_shape(
            1, Inches(0.5), Inches(0.4), Inches(12.333), Inches(1.2)
        )
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = RGBColor(26, 54, 93)
        header_shape.line.fill.background()

        tf = header_shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = s_data["title"]
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)

        p2 = tf.add_paragraph()
        p2.text = s_data["subtitle"]
        p2.font.size = Pt(14)
        p2.font.italic = True
        p2.font.color.rgb = RGBColor(203, 213, 225)

        content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.2))
        ctf = content_box.text_frame
        ctf.word_wrap = True

        for idx, bullet in enumerate(s_data["bullets"]):
            bp = ctf.add_paragraph() if idx > 0 else ctf.paragraphs[0]
            bp.text = f"•  {bullet}"
            bp.font.size = Pt(18)
            bp.font.color.rgb = RGBColor(45, 55, 72)
            bp.space_after = Pt(14)

    prs.save(pptx_path)
    print(f"[SUCCESS] PPTX generated at {pptx_path}")


if __name__ == "__main__":
    generate_pdf()
    generate_pptx()
